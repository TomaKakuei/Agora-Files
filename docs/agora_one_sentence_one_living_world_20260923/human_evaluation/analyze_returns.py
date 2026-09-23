#!/usr/bin/env python3
"""Read the returned PDF fields, preserve exclusions, decode the frozen allocation.

Run with Python 3.9+, pypdf, openpyxl, numpy and matplotlib. No model calls.
The archived returns are input data; this script never edits their answers.
"""
import csv
import io
import json
import logging
import statistics
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
LABELS = {
    'gemini_2_5_flash': 'Gemini 2.5 Flash',
    'gemini_3_1_flash_lite': 'Gemini 3.1 Flash Lite',
    'gemini_3_1_pro_preview': 'Gemini 3.1 Pro Preview',
    'gemini_3_5_flash_lite': 'Gemini 3.5 Flash Lite',
    'gemini_3_7_flash': 'Gemini 3.7 Flash',
    'gpt_5_6_sol': 'GPT-5.6 Sol', 'gpt_5_6_terra': 'GPT-5.6 Terra',
    'agora_compositional': 'Full Agora', 'single_pass': 'Single-pass baseline',
}
ITEMS = {'images': ['Explore preference', 'Map readability', 'Character-world fit'],
         'structure': ['Explore preference', 'World clarity', 'Place-scene fit']}
CHOICES = {'A', 'B', 'Same', 'Not sure'}


def save_csv(name, rows):
    with (HERE/name).open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def summarize(rows, first):
    counts = Counter(r['winner'] for r in rows)
    win = counts[first]
    tie, unsure = counts['Same'], counts['Not sure']
    loss = len(rows)-win-tie-unsure
    denominator = win+tie+loss
    return dict(n=len(rows), wins=win, ties=tie, losses=loss, unsure=unsure,
                judgeable=denominator,
                preference=(win+.5*tie)/denominator if denominator else None)


def extract(d):
    participants, rows, all_fields = [], [], {}
    diagnostics = []
    class Capture(logging.Handler):
        def emit(self, record):
            diagnostics.append(record.getMessage())
    logger = logging.getLogger('pypdf'); handler = Capture()
    logger.addHandler(handler); logger.propagate = False
    with zipfile.ZipFile(HERE/'Survey_DONE.zip') as z:
        assert set(z.namelist()) == {f'Survey_{code}.pdf' for code in d['schedule']}
        for code, tasks in d['schedule'].items():
            reader = PdfReader(io.BytesIO(z.read(f'Survey_{code}.pdf')))
            fields = reader.get_fields()
            expected = {'adult', 'english', 'consent', 'prior_exposure'} | {
                f'R{r}_Q{q}' for r in range(1, 7) for q in range(1, 4)}
            assert set(fields) == expected and len(reader.pages) == 7
            assert code in reader.pages[0].extract_text()
            values = {k: str(v.get('/V', '')).lstrip('/') for k, v in fields.items()}
            # Check the visible widget state independently of the parent /V.
            for key, field in fields.items():
                ref = field.indirect_reference.get_object()
                if '/Kids' in ref:
                    active = [str(k.get_object().get('/AS', '/Off')).lstrip('/')
                              for k in ref['/Kids'] if k.get_object().get('/AS', '/Off') != '/Off']
                    assert active == ([values[key]] if values[key] not in ['', 'Off'] else []), (code, key, active, values[key])
                else:
                    assert str(ref.get('/AS', '/Off')).lstrip('/') == values[key], (code, key)
            all_fields[code] = values
            cover = {k: values[k] if values[k] not in ['', 'Off'] else ''
                     for k in ['adult', 'english', 'consent', 'prior_exposure']}
            valid = sum(values[f'R{r}_Q{q}'] in CHOICES for r in range(1, 7) for q in range(1, 4))
            if any(cover[k] != 'Yes' for k in ['adult', 'english', 'consent']) or not cover['prior_exposure']:
                status = 'Check cover answers'
            elif cover['prior_exposure'] != 'No':
                status = 'Exclude: prior exposure'
            elif valid != 18:
                status = 'Incomplete choices'
            else:
                status = 'Include'
            participants.append(dict(code=code, returned='Yes', **cover, valid_choices=valid, status=status))
            for ri, task in enumerate(tasks, 1):
                pair = d['pairs'][task['pair_id']]
                a = d['artifacts'][task['left']]['condition']
                b = d['artifacts'][task['right']]['condition']
                assert a != b
                for qi in range(1, 4):
                    answer = values[f'R{ri}_Q{qi}']
                    assert answer in CHOICES, (code, ri, qi, answer)
                    winner = {'A': a, 'B': b, 'Same': 'Same', 'Not sure': 'Not sure'}[answer]
                    rows.append(dict(code=code, round=ri, page=ri+1, kind=pair['kind'],
                        pair_id=task['pair_id'], source_id=pair['source_id'], q=qi,
                        item=ITEMS[pair['kind']][qi-1], answer=answer, a=a, b=b,
                        winner=winner, included=status == 'Include'))
    logger.removeHandler(handler)
    unexpected = [m for m in diagnostics if 'Multiple definitions in dictionary' not in m or 'key /FT' not in m]
    assert not unexpected, unexpected
    # The returned editor repeats identical /FT /Btn entries. Values and all
    # appearance states are checked above; never suppress other parser errors.
    return participants, rows, all_fields, len(diagnostics)


def workbook(participants, rows):
    w = openpyxl.load_workbook(HERE/'Results_Template.xlsx')
    w.properties.title = 'Agora visual survey: returned responses, 2026-09-23'
    w.properties.description = 'Actual PDF responses. V04 cover confirmations missing; excluded pending confirmation.'
    pmap = {p['code']: p for p in participants}
    lookup = {(r['code'], r['round'], r['q']): r for r in rows}
    for er in range(7, 18):
        s = w['Participants']; p = pmap[s.cell(er, 1).value]
        for col, key in enumerate(['returned', 'adult', 'english', 'consent', 'prior_exposure'], 2):
            s.cell(er, col).value = p[key] or None
    for er in range(7, 73):
        s = w['Responses']; code, rnd = s.cell(er, 1).value, s.cell(er, 2).value
        for qi in range(1, 4):
            s.cell(er, qi+4).value = lookup[(code, rnd, qi)]['answer']
    s = w['Start here']
    s['B2'] = 'Visual survey | Completed results'
    s['B19'] = '11 returns; 10 included. V04 cover blank: pending confirmation, not imputed.'
    # Replace blank-template instructions with accurate completed-file notes.
    for row in s:
        for cell in row:
            if isinstance(cell.value, str) and 'No human results are included in this template.' in cell.value:
                cell.value = cell.value.replace('No human results are included in this template.', 'This copy contains actual returned choices; the original blank template is retained separately.')
    w.save(HERE/'Results_Completed.xlsx')


def write_tables(d, pair_results, structure_overall):
    generated = PAPER/'sections/generated'; generated.mkdir(exist_ok=True)
    lines = [r'\begin{tabularx}{\columnwidth}{@{}Yrrrrr@{}}', r'\toprule',
             r'Item & Full & Base & Tie & ? & Score \\', r'\midrule']
    for r in structure_overall:
        item = ['Explore', 'Clarity', 'Scene fit'][r['q']-1]
        lines.append(f"{item} & {r['wins']} & {r['losses']} & {r['ties']} & {r['unsure']} & {100*r['equal_request_preference']:.1f}\\% " + r'\\')
    lines += [r'\bottomrule', r'\end{tabularx}']
    (generated/'human_structure_summary.tex').write_text('\n'.join(lines)+'\n')
    lines = [r'\begin{tabularx}{\textwidth}{@{}Yrrr@{}}', r'\toprule',
             r'Request & Explore & Clarity & Scene fit \\', r'\midrule']
    for pid, pair in sorted(d['pairs'].items(), key=lambda kv: kv[1]['source_id']):
        if pair['kind'] != 'structure': continue
        rs = [r for r in pair_results if r['pair_id'] == pid]
        cells = [f"{100*r['preference']:.0f}\\% ({r['wins']}/{r['ties']}/{r['losses']}/{r['unsure']})" for r in rs]
        lines.append(pair['source_id'].replace('_', ' ').title()+' & '+' & '.join(cells)+r' \\')
    lines += [r'\bottomrule', r'\end{tabularx}']
    (generated/'human_structure_detail.tex').write_text('\n'.join(lines)+'\n')
    models = sorted({r['first'] for r in pair_results if r['kind']=='images'} | {r['second'] for r in pair_results if r['kind']=='images'})
    ids = {m: f'M{i+1}' for i, m in enumerate(models)}
    lines = [r'\begin{tabular}{@{}lrrr@{}}', r'\toprule',
             r'Pair & Explore & Map & Characters \\', r'\midrule']
    for pid in sorted({r['pair_id'] for r in pair_results if r['kind']=='images'},
                      key=lambda pid: next((r['first'], r['second']) for r in pair_results if r['pair_id']==pid)):
        rs = [r for r in pair_results if r['pair_id']==pid]
        counts = [f"{r['wins']}/{r['ties']}/{r['losses']}/{r['unsure']}" for r in rs]
        lines.append(ids[rs[0]['first']]+'--'+ids[rs[0]['second']]+' & '+' & '.join(counts)+r' \\')
    lines += [r'\bottomrule', r'\end{tabular}']
    (generated/'human_image_detail.tex').write_text('\n'.join(lines)+'\n')
    (generated/'human_model_key.tex').write_text('; '.join(ids[m]+': '+LABELS[m] for m in models)+'.\n')


def figures(d, pair_results):
    plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':9, 'axes.spines.top':False,
                         'axes.spines.right':False, 'pdf.fonttype':42, 'ps.fonttype':42})
    models = sorted({r['first'] for r in pair_results if r['kind']=='images'} | {r['second'] for r in pair_results if r['kind']=='images'})
    short = ['G2.5 Flash','G3.1 Lite','G3.1 Pro','G3.5 Lite','G3.7 Flash','GPT Sol','GPT Terra']
    fig, grid = plt.subplots(2, 2, figsize=(7.2, 6.5), layout='constrained')
    axes = list(grid.flat[:3])
    for qi, ax in enumerate(axes, 1):
        a = np.full((7, 7), np.nan); annotations = {}
        for r in pair_results:
            if r['kind']!='images' or r['q']!=qi: continue
            i,j = models.index(r['first']),models.index(r['second'])
            if r['preference'] is not None:
                a[i,j]=r['preference']; a[j,i]=1-r['preference']
                annotations[i,j]=annotations[j,i]=r['judgeable']
        ax.set_facecolor('#eef1f3'); im=ax.imshow(a,vmin=0,vmax=1,cmap='BrBG')
        for (i,j),n in annotations.items():
            ax.text(j,i,f'{100*a[i,j]:.0f}\n(n={n})',ha='center',va='center',fontsize=6.6,
                    color='white' if a[i,j]<.16 or a[i,j]>.84 else '#20272b')
        ax.set_xticks(range(7),short,rotation=50,ha='right',fontsize=7)
        ax.set_yticks(range(7),short,fontsize=7); ax.set_title(ITEMS['images'][qi-1],fontweight='bold',pad=12)
        ax.tick_params(length=0)
    key = grid.flat[3]; key.axis('off')
    key.text(0, .97, 'How to read the matrices', va='top', fontsize=10, fontweight='bold')
    key.text(0, .83, 'Cell = preference for the row artifact\nagainst the column artifact.\nSame = half a vote; Not sure is excluded.\nn = judgeable votes for that pair/item.\nMirrored cells reuse the same votes.', va='top', fontsize=8, linespacing=1.5)
    key.text(0, .44, '10 eligible returns | 40 image comparisons\n21 pairs | one shared promise-city scene\nPair coverage: 4 x one, 16 x two, 1 x four\nRecorded identifiers; no global ranking.', va='top', fontsize=8, linespacing=1.5)
    cb = key.inset_axes([0, .045, .9, .065])
    fig.colorbar(im, cax=cb, orientation='horizontal', ticks=[0,.5,1])
    cb.set_xticklabels(['0%', '50%', '100%'], fontsize=8)
    for ext in ['pdf','png']: fig.savefig(PAPER/'figures'/f'human_image_preferences.{ext}',dpi=200,bbox_inches='tight')
    plt.close(fig)
    structure = sorted([p for p in d['pairs'] if d['pairs'][p]['kind']=='structure'],key=lambda p:d['pairs'][p]['source_id'])
    fig,ax=plt.subplots(figsize=(8.5,4.4),layout='constrained')
    for qi,(color,marker) in enumerate([('#156b70','o'),('#a56428','s'),('#715e94','D')],1):
        rs=[next(r for r in pair_results if r['pair_id']==p and r['q']==qi) for p in structure]
        yy=np.arange(10)+(qi-2)*.21
        ax.scatter([r['preference'] for r in rs],yy,label=ITEMS['structure'][qi-1],color=color,marker=marker,s=42,zorder=3)
    ax.axvline(.5,ls='--',color='#9ca5a8',lw=.8)
    ax.set_yticks(range(10),[d['pairs'][p]['source_id'].replace('_',' ').title() for p in structure],fontsize=9)
    ax.invert_yaxis();ax.set_xlim(-.055,1.055);ax.set_xticks([0,.25,.5,.75,1],['0%','25%','50%','75%','100%'])
    ax.set_xlabel('Preference for full Agora within each fixed request (ties = half)')
    ax.grid(axis='x',alpha=.13);ax.legend(loc='upper center',bbox_to_anchor=(.5,1.14),ncol=3,frameon=False)
    for ext in ['pdf','png']:fig.savefig(PAPER/'figures'/f'human_structure_preferences.{ext}',dpi=200,bbox_inches='tight')
    plt.close(fig)


def main():
    d=json.loads((HERE/'dataset.json').read_text())
    participants,rows,fields,parser_notices=extract(d)
    included=[r for r in rows if r['included']]
    save_csv('participants.csv',participants);save_csv('responses_long.csv',rows)
    (HERE/'extracted_fields.json').write_text(json.dumps(fields,indent=2)+'\n')
    pairs=[]
    for pid,pair in sorted(d['pairs'].items()):
        first,second=sorted(d['artifacts'][a]['condition'] for a in pair['candidates'])
        for qi in range(1,4):
            subset=[r for r in included if r['pair_id']==pid and r['q']==qi]
            pairs.append(dict(pair_id=pid,kind=pair['kind'],source_id=pair['source_id'],q=qi,
                              first=first,second=second,**summarize(subset,first)))
    save_csv('pair_results.csv',pairs)
    overall=[]
    for qi in range(1,4):
        rs=[r for r in included if r['kind']=='structure' and r['q']==qi]
        ps=[r['preference'] for r in pairs if r['kind']=='structure' and r['q']==qi]
        overall.append(dict(q=qi,item=ITEMS['structure'][qi-1],**summarize(rs,'agora_compositional'),
                            covered_requests=sum(p is not None for p in ps),
                            equal_request_preference=statistics.mean(ps) if all(p is not None for p in ps) else None))
    image_coverage=Counter(r['n'] for r in pairs if r['kind']=='images' and r['q']==1)
    structure_coverage=Counter(r['n'] for r in pairs if r['kind']=='structure' and r['q']==1)
    side_counts={kind:{str(q):dict(Counter(r['answer'] for r in included if r['kind']==kind and r['q']==q)) for q in range(1,4)} for kind in ITEMS}
    # Descriptive leave-one-participant-out sensitivity; not confidence intervals.
    # Keep the common subset with >=2 raters so every deletion retains coverage.
    common={r['pair_id'] for r in pairs if r['kind']=='structure' and r['q']==1 and r['n']>=2}
    sensitivity=[]
    codes=[p['code'] for p in participants if p['status']=='Include']
    for qi in range(1,4):
        vals=[]
        for omit in codes:
            means=[summarize([r for r in included if r['kind']=='structure' and r['q']==qi and r['pair_id']==pid and r['code']!=omit], 'agora_compositional')['preference'] for pid in common]
            if all(m is not None for m in means):vals.append(statistics.mean(means))
        sensitivity.append(dict(q=qi,common_requests=len(common),valid_deletions=len(vals),
                                minimum=min(vals) if vals else None,maximum=max(vals) if vals else None))
    summary=dict(returned=len(participants),included=len(codes),pending=[p['code'] for p in participants if p['status']!='Include'],
        administration={'recruitment':'Online volunteers, organizer report','compensation':'None, organizer report','V04':'Organizer reports forgotten cover; participant confirmations still missing','ethics_review':'No approval or exemption record supplied','completion_times':'Not recorded'},
        returned_choices=len(rows),included_choices=len(included),image_comparisons=sum(r['kind']=='images' and r['q']==1 for r in included),
        structure_comparisons=sum(r['kind']=='structure' and r['q']==1 for r in included),
        image_pair_coverage=dict(image_coverage),structure_pair_coverage=dict(structure_coverage),
        structure=overall,side_choices=side_counts,leave_one_participant_out_common_requests=sensitivity,
        pdf_field_values_match_widget_states=True,identical_duplicate_field_type_notices=parser_notices,
        limitations=['One visual prompt; fixed outputs and shared image generator.',
                     'The instrument measures artifact preference, not gameplay or full topology.',
                     'One pending cover; do not infer consent from completed choices.',
                     'Organizer reports online unpaid volunteers and an omitted V04 cover; timing, demographics, and ethics review not supplied.',
                     'Model strings are archived backend identifiers, not independently verified product names.',
                     'No population confidence intervals or significance claims; repeated judgments share raters and artifacts.'])
    (HERE/'summary.json').write_text(json.dumps(summary,indent=2)+'\n')
    workbook(participants,rows);write_tables(d,pairs,overall);figures(d,pairs)
    print(json.dumps(summary,indent=2))


if __name__=='__main__':main()
