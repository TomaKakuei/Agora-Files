#!/usr/bin/env python3
"""Descriptive, post-return model summaries; preserve all exact-pair results.

Scores pool observed votes, so opponent weights are not equal. These are not
Elo scores, a composite ranking, or population performance estimates.
"""
import csv
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE=Path(__file__).resolve().parent
LABELS={
    'gemini_2_5_flash':'Gemini 2.5 Flash',
    'gemini_3_1_flash_lite':'Gemini 3.1 Flash Lite',
    'gemini_3_1_pro_preview':'Gemini 3.1 Pro Preview',
    'gemini_3_5_flash_lite':'Gemini 3.5 Flash Lite',
    'gemini_3_7_flash':'Gemini 3.7 Flash',
    'gpt_5_6_sol':'GPT-5.6 Sol', 'gpt_5_6_terra':'GPT-5.6 Terra',
}
ITEMS={1:'Explore preference',2:'Map readability',3:'Character-world fit'}


def main():
    with (HERE/'responses_long.csv').open() as f:
        source=[r for r in csv.DictReader(f) if r['included']=='True' and r['kind']=='images']
    assert len(source)==120 and len({r['code'] for r in source})==10
    rows=[]
    for model in sorted(LABELS):
        for q in range(1,4):
            votes=[r for r in source if int(r['q'])==q and model in [r['a'],r['b']]]
            c=Counter('wins' if r['winner']==model else 'ties' if r['winner']=='Same' else 'unsure' if r['winner']=='Not sure' else 'losses' for r in votes)
            n=c['wins']+c['ties']+c['losses']
            points=c['wins']+.5*c['ties']
            rows.append(dict(model=model,label=LABELS[model],q=q,item=ITEMS[q],
                             comparisons=len(votes),wins=c['wins'],ties=c['ties'],
                             losses=c['losses'],unsure=c['unsure'],judgeable=n,
                             preference_points=points,observed_vote_score=points/n))
    # Independent accounting: each comparison appears in two model summaries.
    for q in range(1,4):
        subset=[r for r in rows if r['q']==q]
        assert sum(r['comparisons'] for r in subset)==80
        assert sum(r['wins'] for r in subset)==sum(r['losses'] for r in subset)
        assert sum(r['judgeable'] for r in subset)==2*sum(r['preference_points'] for r in subset)
    with (HERE/'model_summary.csv').open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]),lineterminator='\n')
        writer.writeheader();writer.writerows(rows)
    tex=[r'\begin{tabularx}{\textwidth}{@{}Yrrr@{}}',r'\toprule',
         r'Recorded configuration & Explore & Map readability & Character fit \\',r'\midrule']
    for model in sorted(LABELS):
        cells=[]
        for r in (r for r in rows if r['model']==model):
            points=f"{r['preference_points']:g}"
            cells.append(f"{points}/{r['judgeable']} ({r['observed_vote_score']*100:.1f}\\%)")
        tex.append(LABELS[model]+' & '+' & '.join(cells)+r' \\')
    tex += [r'\bottomrule',r'\end{tabularx}']
    (HERE.parent/'sections/generated/human_model_summary.tex').write_text('\n'.join(tex)+'\n')
    w=openpyxl.load_workbook(HERE/'Results_Completed.xlsx',data_only=True)
    # A separate values-only workbook leaves the audited original formulas intact.
    out=openpyxl.Workbook(); s=out.active;s.title='Model summary'
    s.append(['Observed model comparisons | one fixed scene'])
    s.append(['10 eligible returns; 40 image comparisons; three separate questions.'])
    s.append(['Score = (wins + 0.5 * ties) / judgeable votes; Not sure is excluded.'])
    s.append(['Post-return descriptive summary. Opponent weights differ; not a general model ranking.'])
    s.append(['Model','Question','Comparisons','Wins','Same','Losses','Not sure','Judgeable','Preference points','Observed vote score'])
    for r in rows:s.append([r['label'],r['item'],r['comparisons'],r['wins'],r['ties'],r['losses'],r['unsure'],r['judgeable'],r['preference_points'],r['observed_vote_score']])
    s.freeze_panes='C6';s.auto_filter.ref='A5:J26'
    s.column_dimensions['A'].width=29;s.column_dimensions['B'].width=26
    for col in 'CDEFGHIJ':s.column_dimensions[col].width=17
    for row in range(1,5):s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=10);s.row_dimensions[row].height=24
    for cell in s[5]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='176B70');cell.alignment=Alignment(wrap_text=True)
    for er in range(6,27):s.cell(er,10).number_format='0.0%'
    s.row_dimensions[5].height=32
    # Preserve all exact image pair results alongside the descriptive roll-up.
    t=out.create_sheet('Exact image pairs')
    for row in w['Image results'].iter_rows():t.append([c.value for c in row])
    t.freeze_panes='D7'
    for col in 'ABC':t.column_dimensions[col].width=29
    t.column_dimensions['D'].width=24
    for col in 'EFGHIJKL':t.column_dimensions[col].width=15
    for er in range(7,70):t.cell(er,12).number_format='0.0%'
    out.save(HERE/'Model_Comparison_Results.xlsx')
    reread=openpyxl.load_workbook(HERE/'Model_Comparison_Results.xlsx',data_only=True)
    for er,r in enumerate(rows,6):
        assert reread['Model summary'].cell(er,10).value==r['observed_vote_score'] or abs(reread['Model summary'].cell(er,10).value-r['observed_vote_score'])<1e-14
    print('Verified 21 model/item summaries against 120 included image choices; exact pair results retained.')


if __name__=='__main__':main()
