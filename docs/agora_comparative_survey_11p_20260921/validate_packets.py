#!/usr/bin/env python3
"""Check frozen evidence, PDF pagination, blinding and real spreadsheet formulas.

Synthetic workbooks are written only under a temporary directory and deleted.
LibreOffice recalculates the test files so assertions inspect actual calculated
values rather than formulas mirrored by the test itself.
"""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

from build_packets import HERE, SOURCE, VERSION, allocate


def norm(text):
    return ''.join(unicodedata.normalize('NFKC',text).split())


def check_materials():
    key=json.loads((HERE/'researcher_only/blind_key.json').read_text())
    data=json.loads((HERE/'researcher_only/frozen_display.json').read_text())
    source_key=json.loads((SOURCE/'blind_key.json').read_text())
    source=json.loads((SOURCE/'public_payload.json').read_text())
    assert data['artifacts']==source['artifacts'] and data['pairs']==source['pairs']
    regenerated,metadata=allocate(source,source_key)
    assert regenerated==key['schedule']==data['schedule'] and metadata==key['allocation']
    assert key['version']==data['version']==VERSION
    for name,expected in json.loads((HERE/'researcher_only/source_hashes.json').read_text()).items():
        assert hashlib.sha256((SOURCE/name).read_bytes()).hexdigest()==expected
    counts=Counter();orientations=Counter();model_c1_a=0;word_total=0
    for code,tasks in key['schedule'].items():
        assert len(tasks)==6 and len({t['pair_id'] for t in tasks})==6
        assert Counter(key['pairs'][t['pair_id']]['arm'] for t in tasks)=={'model':4,'architecture':2}
        assert Counter(key['pairs'][t['pair_id']]['prompt_id'] for t in tasks if key['pairs'][t['pair_id']]['arm']=='model')=={'lwb_hidden_v2_01':2,'lwb_hidden_v2_02':1,'lwb_hidden_v2_03':1}
        arch_a=[]
        with zipfile.ZipFile(HERE/'participants'/f'{code}_问卷.docx') as z:
            assert z.testzip() is None
            xml=z.read('word/document.xml')
        root=ET.fromstring(xml);ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        texts=[e.text or '' for e in root.findall('.//w:t',ns)]
        pdf=HERE/'participants'/f'{code}_问卷.pdf'
        pages=subprocess.check_output(['pdftotext','-raw',str(pdf),'-']).decode().split('\f')
        if not pages[-1].strip():pages.pop()
        assert len(pages)==19,(code,len(pages))
        joined='\n'.join(pages)
        for forbidden in ['gemini_','gpt_5_6','gemini-','gpt-5.6','single_pass','agora_compositional','blind_key','source_records']:
            assert forbidden not in joined.lower() and forbidden not in xml.decode().lower(),(code,forbidden)
        for position,t in enumerate(tasks):
            pid=t['pair_id'];meta=key['pairs'][pid]
            counts[pid]+=1;orientations[(pid,t['left'])]+=1
            a=key['artifacts'][t['left']]['condition'];b=key['artifacts'][t['right']]['condition']
            if meta['arm']=='architecture':arch_a.append(a=='agora_compositional')
            else:model_c1_a+=a==min(a,b)
            assert set([t['left'],t['right']])==set(meta['candidates'])
            for side,offset,label in [('left',1,'A'),('right',2,'B')]:
                page=pages[3*position+offset]
                assert f'第{position+1}/6场' in norm(page) and f'设计{label}' in norm(page)
                assert norm(data['pairs'][pid]['prompt']) in norm(page),(code,position,'prompt')
                for section in data['artifacts'][t[side]]['sections']:
                    for text in section['rows']:
                        assert '• '+text in texts,(code,'DOCX missing evidence')
                        assert norm(text) in norm(page),(code,position,label,'PDF missing evidence',text[:70])
            assert f'第{position+1}/6场' in norm(pages[3*position+3]) and '评分' in pages[3*position+3]
        assert sorted(arch_a)==[False,True]
        # Check all extracted word boxes, including the longest evidence page.
        bbox=subprocess.check_output(['pdftotext','-bbox',str(pdf),'-'])
        bx=ET.fromstring(bbox)
        for page in bx.iter():
            if page.tag.endswith('page'):
                width=float(page.attrib['width']);height=float(page.attrib['height'])
                for word in page.iter():
                    if not word.tag.endswith('word'):continue
                    word_total+=1
                    assert 0<=float(word.attrib['xMin'])<=float(word.attrib['xMax'])<=width+.2
                    assert 0<=float(word.attrib['yMin'])<=float(word.attrib['yMax'])<=height+.2
    assert model_c1_a==22
    model=[p for p,m in key['pairs'].items() if m['arm']=='model']
    arch=[p for p,m in key['pairs'].items() if m['arm']=='architecture']
    assert Counter(counts[p] for p in model)=={1:38,2:3}
    assert Counter(counts[p] for p in arch)=={2:9,4:1}
    for pid in arch+[p for p in model if counts[p]==2]:
        a,b=key['pairs'][pid]['candidates'];assert orientations[pid,a]==orientations[pid,b]
    workbook=openpyxl.load_workbook(HERE/'Agora_11人_结果录入.xlsx')
    for row in range(7,18):
        for col in list(range(2,10))+[12]:assert workbook['参与者登记'].cell(row,col).value is None
    for row in range(7,73):
        for col in range(4,12):assert workbook['评分录入'].cell(row,col).value is None
    assert len(workbook['评分录入'].data_validations.dataValidation)==2
    assert not workbook['评分录入']['D7'].protection.locked
    assert workbook['评分录入']['T7'].protection.locked
    assert workbook['评分录入'].protection.sheet
    return {'packets':11,'pages_each':19,'model_ratings':44,'architecture_ratings':22,'exact_model_pairs':41,'architecture_requests':10,'word_boxes_checked':word_total}


def check_calculation():
    original=HERE/'Agora_11人_结果录入.xlsx'
    key=json.loads((HERE/'researcher_only/blind_key.json').read_text())
    with tempfile.TemporaryDirectory(prefix='agora_11p_SYNTHETIC_') as temp:
        root=Path(temp);source=root/'inputs';output=root/'calculated';source.mkdir();output.mkdir()
        cases=['blank','condition1','condition2','ties_unable','exclusions','all_unable','macro_missing','ordinal_strings']
        for case in cases:
            book=openpyxl.load_workbook(original)
            if case!='blank':
                for row in range(7,18):
                    for col in range(2,6):book['参与者登记'].cell(row,col,'是')
                    book['参与者登记'].cell(row,6,'没有')
                for row in range(7,73):
                    ws=book['评分录入'];c1_on_a=ws.cell(row,16).value==ws.cell(row,18).value
                    winner=2 if c1_on_a else 5
                    for col in range(4,10):
                        value=winner
                        if case=='condition2':value=5 if c1_on_a else 1
                        if case=='ties_unable':value=3 if col!=5 else 6
                        if case=='all_unable':value=6
                        if case=='ordinal_strings':value=str((row-7)%6+1)
                        ws.cell(row,col,value)
                if case=='exclusions':
                    reg=book['参与者登记'];reg['F7']='有';reg['F8']='不确定';reg['E9']='否';reg['D10']=None
                    book['评分录入']['D31']=7  # Q05 invalid score; rejected from the primary population.
                    book['评分录入']['D37']=None  # Q06 incomplete score.
                    reg['B13']='否'  # Q07 not returned.
                if case=='macro_missing':
                    # Remove all primary votes for a whole fixed request; other
                    # items stay valid. Full architecture mean must be blank.
                    target=next(t for ts in key['schedule'].values() for t in ts if key['pairs'][t['pair_id']]['arm']=='architecture')['pair_id']
                    for i,ts in enumerate(key['schedule'].values()):
                        for j,t in enumerate(ts):
                            if t['pair_id']==target:book['评分录入'].cell(7+6*i+j,4,6)
            book.save(source/f'{case}.xlsx')
        command=['libreoffice',f'-env:UserInstallation={(root/"office_profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(output)]+[str(p) for p in sorted(source.glob('*.xlsx'))]
        subprocess.run(command,check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        assert len(list(output.glob('*.xlsx')))==len(cases)
        for case in cases:
            b=openpyxl.load_workbook(output/f'{case}.xlsx',data_only=True)
            for ws in b:
                for row in ws:
                    for cell in row:assert cell.data_type!='e',(case,ws.title,cell.coordinate,cell.value)
            included=sum(b['参与者登记'].cell(r,11).value=='纳入' for r in range(7,18))
            assert included==(0 if case=='blank' else 4 if case=='exclusions' else 11),(case,included)
            rows=list(b['逐题汇总'].iter_rows(min_row=7,max_row=312,values_only=True))
            for row in rows:
                arm,prompt,c1,c2,item,planned,total,valid,wins,ties,losses,unable,mean=row[:13]
                assert total==valid+unable and valid==wins+ties+losses,(case,row)
                if case=='blank':assert total==0 and mean in (None,'')
                if case=='condition1':assert total==planned and wins==valid==planned and mean==1
                if case=='condition2':assert total==planned and losses==valid==planned and mean==0
                if case=='all_unable':assert unable==planned and valid==0 and mean in (None,'')
                if case=='ties_unable':
                    if item=='设定实现':assert unable==planned and mean in (None,'')
                    else:assert ties==valid==planned and mean==.5
            macro=b['架构汇总']['I19'].value
            if case in ['blank','all_unable','macro_missing']:assert macro in (None,''),(case,macro)
            if case=='condition1':assert macro==1
            if case=='condition2':assert macro==0
            if case=='ties_unable':assert macro==.5
            if case=='macro_missing':assert b['架构汇总']['I20'].value==9
            if case=='ordinal_strings':
                for row in range(7,73):
                    ws=b['评分录入'];raw=int(ws.cell(row,4).value);score=ws.cell(row,20).value
                    left_is_first=ws.cell(row,16).value==ws.cell(row,18).value
                    expected='' if raw==6 else .5 if raw==3 else int((raw<3)==left_is_first)
                    assert score==expected or (expected=='' and score is None),(row,raw,score,expected)
                # Independent calculation of prompt-weighted architecture mean
                # from the synthetic raw scores, not from its summary formulas.
                prompts={}
                for r in range(7,73):
                    ws=b['评分录入']
                    if ws.cell(r,14).value!='架构':continue
                    raw=int(ws.cell(r,4).value)
                    if raw==6:continue
                    left_first=ws.cell(r,16).value==ws.cell(r,18).value
                    value=.5 if raw==3 else int((raw<3)==left_first)
                    prompts.setdefault(ws.cell(r,15).value,[]).append(value)
                if len(prompts)==10:
                    expected=sum(sum(v)/len(v) for v in prompts.values())/10
                    assert abs(macro-expected)<1e-10
            for r in range(7,28):
                ws=b['模型汇总'];planned=ws.cell(r,3).value;observed=ws.cell(r,4).value;mean=ws.cell(r,9).value
                if observed<planned:assert mean in (None,'')
                if case=='condition1':assert mean==1
                if case=='condition2':assert mean==0
        return {'engine':'LibreOffice Calc','synthetic_cases':cases,'synthetic_files_published':0,'formula_errors':0}


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--skip-calculation',action='store_true');args=ap.parse_args()
    result={'materials':check_materials()}
    if not args.skip_calculation:result['spreadsheet']=check_calculation()
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
