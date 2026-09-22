#!/usr/bin/env python3
"""Validate genuine assets, allocations, blank PDF forms and Calc formulas."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from prepare_materials import HERE,allocate


def check_materials():
    d=json.loads((HERE/'researcher_only/dataset.json').read_text())
    schedule,allocation=allocate(d['pairs']);assert schedule==d['schedule'] and allocation==d['allocation']
    for row in json.loads((HERE/'researcher_only/SOURCE_MANIFEST.json').read_text()):
        p=HERE/row['file'];assert p.stat().st_size==row['bytes'] and hashlib.sha256(p.read_bytes()).hexdigest()==row['sha256']
    for a in d['artifacts'].values():
        spec=json.loads((HERE/a['spec']).read_text())
        if a['kind']=='structure':
            assert len(a['rooms'])==len(spec['rooms'])
            for shown,original in zip(a['rooms'],spec['rooms']):
                assert shown['name']==original['name']
                assert shown['residents']==sum(str(p.get('home_base','')).strip().casefold()==original['name'].strip().casefold() for p in spec['main_characters'])
                props=original.get('scene_components',original.get('visual',{}).get('scene_components',[]))
                assert shown['props']==len(props)
                assert shown['example_prop']==(props[0].get('label','') if props else '')
            assert sum(r['residents'] for r in a['rooms'])+a['unassigned_people']==len(spec['main_characters'])
    frequency=Counter();orientation=Counter();boxes=0;image_counts=[]
    for code,tasks in d['schedule'].items():
        assert len(tasks)==6 and len({t['pair_id'] for t in tasks})==6
        assert Counter(d['pairs'][t['pair_id']]['kind'] for t in tasks)=={'images':4,'structure':2}
        arch=[]
        for t in tasks:
            pair=d['pairs'][t['pair_id']];assert {t['left'],t['right']}==set(pair['candidates'])
            frequency[t['pair_id']]+=1;orientation[(t['pair_id'],t['left'])]+=1
            if pair['kind']=='structure':arch.append(d['artifacts'][t['left']]['condition']=='agora_compositional')
        assert sorted(arch)==[False,True]
        pdf=HERE/'participants'/f'Survey_{code}.pdf'
        info=subprocess.check_output(['pdfinfo',str(pdf)]).decode();assert re.search(r'Pages:\s+7\b',info) and 'AcroForm' in info
        raw=pdf.read_bytes();fields=re.findall(rb'/T \(([^)]*)\)',raw)
        expected=['adult','english','consent','prior_exposure']+[f'R{i}_Q{j}' for i in range(1,7) for j in range(1,4)]
        assert set(x.decode() for x in fields)==set(expected) and len(fields)==22
        assert all(v==b'Off' for v in re.findall(rb'/V /([^\s<>]+)',raw))
        pages=subprocess.check_output(['pdftotext','-raw',str(pdf),'-']).decode().split('\f')
        if not pages[-1].strip():pages.pop()
        assert len(pages)==7
        for i,t in enumerate(tasks,1):
            page=pages[i];pair=d['pairs'][t['pair_id']]
            assert ' '.join(pair['scene'].split()) in ' '.join(page.split())
            assert 'Which world would you rather explore?' in page
            assert len(page.split())<300,(code,i,len(page.split()))
        joined='\n'.join(pages)
        assert not re.search('[\u3400-\u9fff]',joined)
        for forbidden in ['gemini','gpt-5','single_pass','agora_compositional','source_records']:
            assert forbidden not in joined.lower(),(code,forbidden)
        tree=ET.fromstring(subprocess.check_output(['pdftotext','-bbox',str(pdf),'-']))
        for p in tree.iter():
            if p.tag.endswith('page'):
                width=float(p.attrib['width']);height=float(p.attrib['height'])
                for word in p.iter():
                    if word.tag.endswith('word'):
                        boxes+=1
                        assert 0<=float(word.attrib['xMin'])<=float(word.attrib['xMax'])<=width+.2
                        assert 0<=float(word.attrib['yMin'])<=float(word.attrib['yMax'])<=height+.2
        with zipfile.ZipFile(HERE/'participants'/f'Survey_{code}.docx') as z:
            assert z.testzip() is None
            xml=z.read('word/document.xml').decode();assert xml.count('>[choose]</w:t>')==18
            assert not re.search('[\u3400-\u9fff]',xml)
            assert len([n for n in z.namelist() if n.startswith('word/media/')])==6
        listing=subprocess.check_output(['pdfimages','-list',str(pdf)]).decode()
        # 7 maps + 14 atlases are reused as XObjects; image counts depend on
        # within-person model recurrence, so only require actual bitmap content.
        assert len(listing.splitlines())>10
    for kind in ['images','structure']:
        counts=Counter(frequency[p] for p,v in d['pairs'].items() if v['kind']==kind)
        assert counts==({2:20,4:1} if kind=='images' else {2:9,4:1})
    for p,v in d['pairs'].items():
        a,b=v['candidates'];assert orientation[p,a]==orientation[p,b]
    wb=openpyxl.load_workbook(HERE/'Results_Entry.xlsx')
    for row in range(7,18):
        for col in range(2,7):assert wb['Participants'].cell(row,col).value is None
    for row in range(7,73):
        for col in range(5,8):assert wb['Responses'].cell(row,col).value is None
    assert not wb._external_links
    for ws in wb:
        for row in ws:
            for c in row:
                if isinstance(c.value,str):assert not re.search('[\u3400-\u9fff]',c.value)
    return {'packets':11,'pages_per_pdf':7,'fields_per_pdf':22,'choices_per_person':18,'visual_pairs':21,'structure_pairs':10,'word_boxes_checked':boxes,'blank_workbook':True}


def check_office():
    with tempfile.TemporaryDirectory(prefix='agora_VISUAL_SYNTHETIC_') as tmp:
        tmp=Path(tmp);inputs=tmp/'inputs';outputs=tmp/'outputs';inputs.mkdir();outputs.mkdir()
        cases=['blank','source1','source2','same_unsure','excluded','all_unsure','missing_scene']
        for case in cases:
            b=openpyxl.load_workbook(HERE/'Results_Entry.xlsx')
            if case!='blank':
                for r in range(7,18):
                    for col in range(2,6):b['Participants'].cell(r,col,'Yes')
                    b['Participants'].cell(r,6,'No')
                for r in range(7,73):
                    s=b['Responses'];a_first=s.cell(r,12).value==s.cell(r,14).value
                    for col in range(5,8):
                        choice=('A' if a_first else 'B')
                        if case=='source2':choice='B' if a_first else 'A'
                        if case=='same_unsure':choice='Not sure' if col==6 else 'Same'
                        if case=='all_unsure':choice='Not sure'
                        s.cell(r,col,choice)
                if case=='excluded':
                    p=b['Participants'];p['F7']='Yes';p['F8']='Not sure';p['E9']='No';p['D10']=None
                    b['Responses']['E31']='INVALID';b['Responses']['E37']=None;p['B13']='No'
                if case=='missing_scene':
                    e=b['Responses'];target=next(e.cell(r,11).value for r in range(7,73) if e.cell(r,10).value=='structure')
                    for r in range(7,73):
                        if e.cell(r,11).value==target:e.cell(r,5,'Not sure')
            b.save(inputs/(case+'.xlsx'))
        command=['libreoffice',f'-env:UserInstallation={(tmp/"calc_profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(outputs)]+[str(p) for p in inputs.glob('*.xlsx')]
        subprocess.run(command,check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        assert len(list(outputs.glob('*.xlsx')))==len(cases)
        for case in cases:
            b=openpyxl.load_workbook(outputs/(case+'.xlsx'),data_only=True)
            for s in b:
                for row in s:
                    for cell in row:assert cell.data_type!='e',(case,s.title,cell.coordinate,cell.value)
            included=sum(b['Participants'].cell(r,8).value=='Include' for r in range(7,18))
            assert included==(0 if case=='blank' else 4 if case=='excluded' else 11),(case,included)
            for name,last in [('Image results',69),('Structure results',36)]:
                for row in b[name].iter_rows(min_row=7,max_row=last,values_only=True):
                    _,_,_,question,planned,total,judged,wins,same,losses,unsure,mean=row[:12]
                    assert total==judged+unsure and judged==wins+same+losses
                    if case=='blank':assert total==0 and mean in ('',None)
                    if case=='source1':assert judged==wins==planned and mean==1
                    if case=='source2':assert judged==losses==planned and mean==0
                    if case=='all_unsure':assert unsure==planned and mean in ('',None)
                    if case=='same_unsure':
                        if question in ['Map readability','World clarity']:assert unsure==planned and mean in ('',None)
                        else:assert same==judged==planned and mean==.5
            macro=b['Structure results']['L40'].value
            if case=='source1':assert macro==1
            if case=='source2':assert macro==0
            if case=='same_unsure':assert macro==.5
            if case in ['blank','all_unsure','missing_scene']:assert macro in ('',None)
            if case=='missing_scene':assert b['Structure results']['K40'].value==9
        # Word rendering checks use private temporary outputs, never replacing
        # the distributed fillable PDFs with flattened office exports.
        docout=tmp/'word_pdfs';docout.mkdir()
        command=['libreoffice',f'-env:UserInstallation={(tmp/"word_profile").as_uri()}','--headless','--convert-to','pdf','--outdir',str(docout)]+[str(p) for p in sorted((HERE/'participants').glob('*.docx'))]
        subprocess.run(command,check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        assert len(list(docout.glob('*.pdf')))==11
        for p in docout.glob('*.pdf'):
            info=subprocess.check_output(['pdfinfo',str(p)]).decode();assert re.search(r'Pages:\s+7\b',info),(p.name,info)
        return {'synthetic_calculation_cases':len(cases),'formula_errors':0,'word_documents_opened':11,'word_pages_each':7,'synthetic_data_published':0}


if __name__=='__main__':
    result={'materials':check_materials(),'office':check_office()}
    print(json.dumps(result,indent=2))
